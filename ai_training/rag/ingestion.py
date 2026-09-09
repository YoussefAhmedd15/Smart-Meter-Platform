import os
import re
import hashlib
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import time
from concurrent.futures import ThreadPoolExecutor, TimeoutError

import fitz  # PyMuPDF
import openpyxl
import pandas as pd


# ============================================================
# CLASSIFICATION CATEGORIES & PATTERNS
# ============================================================

CATEGORIES = [
    "TECHNICAL_SPECIFICATION",
    "TECHNICAL_DESCRIPTION",
    "INSTALLATION_GUIDE",
    "METER_MANUAL",
    "BROCHURE",
    "ERROR_REFERENCE",
    "SOFTWARE_DOCUMENTATION",
    "L1_SUPPORT",
    "L2_SUPPORT",
    "TEST_DATA",
    "TRACKER",
    "LCD_SPECIFICATION",
    "COMMUNICATION",
    "SMS",
    "FIRMWARE",
    "HARDWARE",
    "CONFIGURATION",
    "HISTORICAL_RECORD",
    "OTHER",
]

METER_MODELS = ["MT174", "MT514", "ME514", "MT880", "MT514-CT", "ME514-5", "SCDC", "MPMS 3000", "Zsk08206"]


def detect_file_hash(file_path: Path) -> str:
    """Calculate SHA-256 hash of a file."""
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def detect_meter_models(text: str, filename: str) -> Tuple[str, str]:
    """
    Detect primary meter model and family from filename and text content.
    Returns (meter_model, meter_family).
    """
    combined = f"{filename} {text[:2000]}".upper()

    if "MT174" in combined:
        return "MT174", "MT"
    elif "MT880" in combined:
        return "MT880", "MT"
    elif "MT514-CT" in combined:
        return "MT514-CT", "MT"
    elif "MT514" in combined:
        return "MT514", "MT"
    elif "ME514-5" in combined or "ME 514-5" in combined:
        return "ME514-5", "ME"
    elif "ME514" in combined or "ME 514" in combined:
        return "ME514", "ME"
    elif "SCDC" in combined or "MPMS" in combined:
        return "SCDC", "SCDC"
    elif "ZSK08206" in combined:
        return "Zsk08206", "LCD"

    return "UNKNOWN", "UNKNOWN"


def detect_document_types(filename: str, text: str) -> List[str]:
    """
    Multi-label classification of document based on filename and extracted text.
    """
    types = set()
    fn_lower = filename.lower()
    text_sample = text[:4000].lower()
    combined = f"{fn_lower} {text_sample}"

    # 1. Trackers & Support Logs
    if "tracker" in fn_lower or ".xlsx" in fn_lower:
        types.add("TRACKER")
        types.add("HISTORICAL_RECORD")
    if "l1 created" in fn_lower or "l1" in fn_lower:
        types.add("L1_SUPPORT")
        types.add("HISTORICAL_RECORD")
    if "l2 soft" in fn_lower or "l2" in fn_lower:
        types.add("L2_SUPPORT")
        types.add("HISTORICAL_RECORD")

    # 2. Installation Guides
    if "installation" in combined or "install" in combined or "guid" in combined:
        types.add("INSTALLATION_GUIDE")
        types.add("HARDWARE")

    # 3. Error Reference
    if "error" in combined or "errors" in combined or "كود" in text_sample or "خطأ" in text_sample:
        types.add("ERROR_REFERENCE")

    # 4. LCD Specifications
    if "lcd" in combined or "meter pages" in combined or "display" in combined:
        types.add("LCD_SPECIFICATION")

    # 5. Technical Specifications / Descriptions
    if "technical specification" in combined or "specs" in combined:
        types.add("TECHNICAL_SPECIFICATION")
    if "technical description" in combined:
        types.add("TECHNICAL_DESCRIPTION")

    # 6. Brochures
    if "brochure" in combined:
        types.add("BROCHURE")

    # 7. Manuals
    if "manual" in combined or "meter.pdf" in fn_lower or "me514.pdf" in fn_lower or "mt174.pdf" in fn_lower:
        types.add("METER_MANUAL")

    # 8. Communication & SMS
    if "sms" in combined:
        types.add("SMS")
        types.add("COMMUNICATION")
    if "communication" in combined or "modem" in combined or "rs485" in combined or "optical" in combined or "gprs" in combined:
        types.add("COMMUNICATION")

    # 9. Software / MPMS / SCDC
    if "mpms" in combined or "scdc" in combined or "software" in combined:
        types.add("SOFTWARE_DOCUMENTATION")

    # 10. Firmware & Hardware
    if "firmware" in combined:
        types.add("FIRMWARE")
    if "terminal" in combined or "wiring" in combined or "hardware" in combined:
        types.add("HARDWARE")

    # 11. Configuration
    if "configuration" in combined or "setting" in combined:
        types.add("CONFIGURATION")

    if not types:
        types.add("OTHER")

    return sorted(list(types))


def detect_version_year(filename: str, text: str) -> Tuple[str, str]:
    """Detect version/revision and year."""
    combined = f"{filename} {text[:3000]}"

    # Year
    year_match = re.search(r"\b(20[12]\d)\b", combined)
    year = year_match.group(1) if year_match else "UNKNOWN"

    # Version / Revision
    ver_match = re.search(r"\b(?:version|revision|rev|ver|v)[.:\s]*([0-9]+(?:\.[0-9]+)*)\b", combined, re.IGNORECASE)
    if ver_match:
        version = f"v{ver_match.group(1)}"
    elif "V5" in combined or "v5" in combined:
        version = "v5"
    elif "4.0" in combined:
        version = "v4.0"
    elif "0-2" in combined or "0.2" in combined:
        version = "v0.2"
    else:
        version = "UNKNOWN"

    return version, year


def detect_language(text: str) -> str:
    """Detect dominant language(s)."""
    has_arabic = bool(re.search(r"[\u0600-\u06FF]", text))
    has_english = bool(re.search(r"[a-zA-Z]", text))

    if has_arabic and has_english:
        return "Arabic / English (Bilingual)"
    elif has_arabic:
        return "Arabic"
    elif has_english:
        return "English"
    return "Multilingual / Technical"


# ============================================================
# PARSERS FOR EACH FILE TYPE
# ============================================================

def parse_pdf(pdf_path: Path) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Extract pages from PDF with page numbers."""
    pages = []
    warnings = []
    try:
        doc = fitz.open(pdf_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()
            if text:
                pages.append({
                    "page": page_num + 1,
                    "text": text,
                })
        doc.close()
    except Exception as e:
        warnings.append(f"PyMuPDF error: {e}")
        # Fallback to pypdf if available
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(pdf_path))
            for page_num, p in enumerate(reader.pages, start=1):
                t = (p.extract_text() or "").strip()
                if t:
                    pages.append({
                        "page": page_num,
                        "text": t,
                    })
        except Exception as e2:
            warnings.append(f"Fallback pypdf error: {e2}")

    return pages, warnings


def parse_pptx(pptx_path: Path) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Extract slides from PPTX via zipfile XML parsing."""
    slides = []
    warnings = []
    try:
        with zipfile.ZipFile(pptx_path) as z:
            slide_files = [f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")]
            def slide_key(name):
                m = re.search(r"slide(\d+)\.xml", name)
                return int(m.group(1)) if m else 999
            slide_files.sort(key=slide_key)

            for idx, sf in enumerate(slide_files, start=1):
                tree = ET.fromstring(z.read(sf))
                texts = [elem.text.strip() for elem in tree.iter() if elem.text and elem.tag.endswith("}t")]
                content = "\n".join([t for t in texts if t])
                if content:
                    slides.append({
                        "page": idx,
                        "text": content,
                    })
    except Exception as e:
        warnings.append(f"PPTX parse error: {e}")

    return slides, warnings


def parse_xlsx(xlsx_path: Path) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Extract structured records from XLSX sheets.
    Batches rows with column headers so context is preserved.
    """
    records = []
    warnings = []
    wb = None
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            headers = []
            rows_data = []
            consecutive_empty = 0

            for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                str_row = [str(c).strip() if c is not None else "" for c in row]
                if not any(str_row):
                    consecutive_empty += 1
                    if consecutive_empty > 1000:
                        break
                    continue
                consecutive_empty = 0

                if not headers:
                    headers = str_row
                    continue

                row_items = []
                for h, v in zip(headers, str_row):
                    if v and v != "None" and h:
                        row_items.append(f"{h}: {v}")

                if row_items:
                    record_text = f"Sheet: {sheet_name} | Row {r_idx} | " + " ; ".join(row_items)
                    rows_data.append((r_idx, record_text))

            batch_size = 5
            for i in range(0, len(rows_data), batch_size):
                batch = rows_data[i:i + batch_size]
                combined_text = "\n".join([item[1] for item in batch])
                start_row = batch[0][0]
                records.append({
                    "page": start_row,
                    "sheet": sheet_name,
                    "text": combined_text,
                })

    except Exception as e:
        warnings.append(f"XLSX parse error: {e}")
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass

    return records, warnings


def parse_csv(csv_path: Path) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Extract structured records from CSV files.
    Batches rows preserving ticket ID, work item type, and state.
    """
    records = []
    warnings = []
    if csv_path.stat().st_size == 0:
        warnings.append("Corrupt/Empty file (0 bytes)")
        return records, warnings

    try:
        df = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
        headers = list(df.columns)

        batch_size = 5
        rows_data = []
        for idx, row in df.iterrows():
            row_num = idx + 2
            row_items = []
            for h in headers:
                val = str(row[h]).strip()
                if val:
                    row_items.append(f"{h}: {val}")
            if row_items:
                rows_data.append((row_num, f"CSV Row {row_num} | " + " ; ".join(row_items)))

        for i in range(0, len(rows_data), batch_size):
            batch = rows_data[i:i + batch_size]
            combined_text = "\n".join([item[1] for item in batch])
            start_row = batch[0][0]
            records.append({
                "page": start_row,
                "sheet": "CSV",
                "text": combined_text,
            })

    except Exception as e:
        warnings.append(f"CSV parse error: {e}")

    return records, warnings


# ============================================================
# MASTER DOCUMENT INVENTORY & PROCESSOR
# ============================================================

class DocumentKnowledgeProcessor:
    """
    Processes all knowledge files, detects duplicates, extracts metadata,
    and produces structured document chunks with full provenance.
    """

    def __init__(self, knowledge_dir: Path):
        self.knowledge_dir = knowledge_dir
        self.inventory: Dict[str, Dict[str, Any]] = {}
        self.seen_hashes: Dict[str, str] = {}
        self.duplicates: List[Dict[str, str]] = []

    def inspect_and_catalog(self) -> Dict[str, Dict[str, Any]]:
        """
        Phase 1 & Phase 2: Catalog every file in knowledge dir.
        Detects exact duplicates and corrupt empty files.
        """
        files = sorted(list(self.knowledge_dir.iterdir()))

        for f in files:
            if f.is_dir():
                continue

            size_bytes = f.stat().st_size
            ext = f.suffix.lower()
            sha256 = detect_file_hash(f)

            is_duplicate = False
            duplicate_of = None

            if size_bytes == 0:
                is_empty = True
                status = "EMPTY_CORRUPT"
            elif sha256 in self.seen_hashes:
                is_duplicate = True
                duplicate_of = self.seen_hashes[sha256]
                status = "EXACT_DUPLICATE"
                self.duplicates.append({
                    "filename": f.name,
                    "duplicate_of": duplicate_of,
                    "sha256": sha256,
                })
            else:
                is_empty = False
                self.seen_hashes[sha256] = f.name
                status = "UNIQUE"

            item: Dict[str, Any] = {
                "filename": f.name,
                "extension": ext,
                "file_type": ext.lstrip("."),
                "size_bytes": size_bytes,
                "sha256": sha256,
                "status": status,
                "is_duplicate": is_duplicate,
                "duplicate_of": duplicate_of,
                "is_empty": (size_bytes == 0),
                "warnings": [],
            }

            self.inventory[f.name] = item

        return self.inventory

    def _process_single_file(self, filename: str, info: Dict[str, Any]) -> List[Dict[str, Any]]:
        """Process a single file: extract records, detect metadata, generate chunks."""
        f_path = self.knowledge_dir / filename
        ext = info["extension"]

        if info.get("is_empty") or f_path.stat().st_size == 0:
            info["is_empty"] = True
            info["document_types"] = ["HISTORICAL_RECORD"]
            info["primary_type"] = "HISTORICAL_RECORD"
            info["meter_model"] = "UNKNOWN"
            info["meter_family"] = "UNKNOWN"
            info["version"] = "UNKNOWN"
            info["year"] = "UNKNOWN"
            info["language"] = "UNKNOWN"
            info["total_pages_or_rows"] = 0
            info["extraction_status"] = "FAILED"
            raise ValueError(f"File is empty (0 bytes): {filename}")

        records: List[Dict[str, Any]] = []
        warnings: List[str] = []

        if ext == ".pdf":
            records, warnings = parse_pdf(f_path)
        elif ext == ".pptx":
            records, warnings = parse_pptx(f_path)
        elif ext == ".xlsx":
            records, warnings = parse_xlsx(f_path)
        elif ext == ".csv":
            records, warnings = parse_csv(f_path)
        else:
            warnings.append(f"Unsupported file extension: {ext}")
            raise ValueError(f"Unsupported file extension: {ext}")

        info["warnings"].extend(warnings)
        info["total_pages_or_rows"] = len(records)

        sample_text = " ".join([r.get("text", "")[:500] for r in records[:5]])

        model, family = detect_meter_models(sample_text, filename)
        doc_types = detect_document_types(filename, sample_text)
        version, year = detect_version_year(filename, sample_text)
        language = detect_language(sample_text)

        info["meter_model"] = model
        info["meter_family"] = family
        info["document_types"] = doc_types
        info["primary_type"] = doc_types[0] if doc_types else "OTHER"
        info["version"] = version
        info["year"] = year
        info["language"] = language
        info["extraction_status"] = "SUCCESS" if records else "FAILED"

        if info["is_duplicate"]:
            info["extraction_status"] = "DUPLICATE_SKIPPED_FOR_INDEX"
            return []

        file_chunks = []
        for r_idx, rec in enumerate(records):
            text = rec.get("text", "").strip()
            if not text:
                continue

            page_num = rec.get("page", r_idx + 1)
            sheet_name = rec.get("sheet", "")

            sub_chunks = self._chunk_text(text, max_chars=1200, overlap=150)
            for c_idx, sub_text in enumerate(sub_chunks):
                clean_sheet = re.sub(r'[^a-zA-Z0-9]', '', sheet_name)
                chunk_id = f"{f_path.stem}_{clean_sheet or 'p'}{page_num}_c{c_idx}"
                file_chunks.append({
                    "id": chunk_id,
                    "text": sub_text,
                    "metadata": {
                        "source": filename,
                        "file_type": ext.lstrip("."),
                        "document_type": info["primary_type"],
                        "document_types": ",".join(doc_types),
                        "meter_model": model,
                        "meter_family": family,
                        "version": version,
                        "year": year,
                        "page": page_num,
                        "sheet": sheet_name,
                        "chunk_index": c_idx,
                        "is_tracker": "TRACKER" in doc_types or "HISTORICAL_RECORD" in doc_types,
                        "status": "VERIFIED_SOURCE",
                    }
                })

        return file_chunks

    def extract_chunks(self, timeout_per_file: int = 120) -> Tuple[List[Dict[str, Any]], Dict[str, Dict[str, Any]]]:
        """
        Phase 3 & Phase 4: Extract all chunks with rich provenance.
        """
        if not self.inventory:
            self.inspect_and_catalog()

        all_chunks: List[Dict[str, Any]] = []
        total_files = len(self.inventory)

        for idx, (filename, info) in enumerate(self.inventory.items(), start=1):
            print(f"[EXTRACTING {idx}/{total_files}] {filename}", flush=True)
            t_start = time.time()

            executor = ThreadPoolExecutor(max_workers=1)
            try:
                future = executor.submit(self._process_single_file, filename, info)
                try:
                    file_chunks = future.result(timeout=timeout_per_file)
                    executor.shutdown(wait=False)
                except TimeoutError:
                    executor.shutdown(wait=False, cancel_futures=True)
                    raise TimeoutError(f"Extraction timed out after {timeout_per_file}s")
                except Exception:
                    executor.shutdown(wait=False)
                    raise

                elapsed = time.time() - t_start
                all_chunks.extend(file_chunks)
                print(f"[EXTRACTED] {filename} | chunks={len(file_chunks)} | time={elapsed:.2f}s", flush=True)

            except Exception as exc:
                elapsed = time.time() - t_start
                print(f"[FAILED] {filename}", flush=True)
                print(f"[ERROR] {exc}", flush=True)
                info["extraction_status"] = "FAILED"
                info["warnings"].append(str(exc))
                continue

        return all_chunks, self.inventory

    def _chunk_text(self, text: str, max_chars: int = 1200, overlap: int = 150) -> List[str]:
        """Split text into character-bounded chunks."""
        text = text.replace("\x00", " ").strip()
        text = re.sub(r"[ \t]+", " ", text)
        if len(text) <= max_chars:
            return [text]

        chunks = []
        start = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            if end < len(text):
                break_pos = text.rfind("\n", start + max_chars // 2, end)
                if break_pos == -1:
                    break_pos = text.rfind(". ", start + max_chars // 2, end)
                if break_pos != -1 and break_pos > start:
                    end = break_pos + 1
            chunks.append(text[start:end].strip())
            if end >= len(text):
                break
            prev_start = start
            start = max(start + 1, end - overlap)
            if start <= prev_start:
                break
        return chunks

